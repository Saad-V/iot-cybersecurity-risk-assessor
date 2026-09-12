import csv
import json
from collections import defaultdict

from google import genai
from docx import Document


# ============================================================
# Security Assessment Functions
# ============================================================

def check_issue(device, metric, issue):
    """Check whether a device has a specific security issue."""

    if metric in ["authentication", "encryption", "logging", "backup"]:
        if device[metric].lower() == "no":
            return issue

    elif metric == "firmware":
        if device[metric].lower() == "outdated":
            return issue

    return None


def get_severity(risk_score):
    """Convert a numerical risk score into a severity level."""

    if 1 <= risk_score <= 5:
        return "Low"

    elif 6 <= risk_score <= 10:
        return "Medium"

    elif 11 <= risk_score <= 15:
        return "High"

    elif 16 <= risk_score <= 25:
        return "Critical"

    return "Invalid"


# ============================================================
# Assessment Configuration
# ============================================================

checks = [
    ("authentication", "Missing authentication", 5, 5, "PR.AA-03"),
    ("encryption", "Missing encryption", 4, 5, "PR.DS-02"),
    ("logging", "Missing security logging", 3, 4, "PR.PS-04"),
    ("backup", "Missing backup", 2, 5, "PR.DS-11"),
    ("firmware", "Outdated firmware", 4, 4, "PR.PS-02"),
]


csf_mapping = {
    "PR.AA-03": {
        "control": "Users, services, and hardware are authenticated",
        "recommendation": (
            "Implement appropriate authentication for users, "
            "services, and devices."
        ),
    },

    "PR.DS-02": {
        "control": (
            "The confidentiality, integrity, and availability "
            "of data-in-transit are protected"
        ),
        "recommendation": (
            "Use appropriate encryption and cryptographic "
            "mechanisms to protect device communications."
        ),
    },

    "PR.PS-04": {
        "control": (
            "Log records are generated and made available "
            "for continuous monitoring"
        ),
        "recommendation": (
            "Configure devices and services to generate security "
            "logs and make them available to a logging system."
        ),
    },

    "PR.DS-11": {
        "control": (
            "Backups of data are created, protected, maintained, "
            "and tested"
        ),
        "recommendation": (
            "Create, protect, maintain, and periodically test backups."
        ),
    },

    "PR.PS-02": {
        "control": (
            "Software is maintained, replaced, and removed "
            "commensurate with risk"
        ),
        "recommendation": (
            "Apply security updates and replace unsupported "
            "firmware/software according to risk."
        ),
    },
}


# ============================================================
# Run Security Assessment
# ============================================================

findings_by_device = defaultdict(list)
total_devices = 0


with open("data/devices.csv", "r") as file:

    devices = csv.DictReader(file)

    for device in devices:

        total_devices += 1

        device_id = device["device_id"]

        for metric, issue, likelihood, impact, csf_control in checks:

            result = check_issue(
                device,
                metric,
                issue
            )

            if result:

                risk_score = likelihood * impact
                severity = get_severity(risk_score)

                finding = {
                    "issue": result,
                    "likelihood": likelihood,
                    "impact": impact,
                    "risk_score": risk_score,
                    "severity": severity,
                    "csf_control": csf_control,
                    "control": csf_mapping[csf_control]["control"],
                    "recommendation": csf_mapping[csf_control]["recommendation"],
                }

                findings_by_device[device_id].append(finding)


# ============================================================
# Create Device Summaries
# ============================================================

device_summaries = {}


for device_id, findings in findings_by_device.items():

    critical_count = sum(
        1
        for finding in findings
        if finding["severity"] == "Critical"
    )

    high_count = sum(
        1
        for finding in findings
        if finding["severity"] == "High"
    )

    medium_count = sum(
        1
        for finding in findings
        if finding["severity"] == "Medium"
    )

    low_count = sum(
        1
        for finding in findings
        if finding["severity"] == "Low"
    )

    highest_risk_finding = max(
        findings,
        key=lambda finding: finding["risk_score"]
    )

    device_summaries[device_id] = {
        "total_findings": len(findings),
        "critical": critical_count,
        "high": high_count,
        "medium": medium_count,
        "low": low_count,
        "highest_risk": highest_risk_finding,
    }


# ============================================================
# Create Overall Summary
# ============================================================

total_findings = sum(
    len(findings)
    for findings in findings_by_device.values()
)

total_critical = sum(
    1
    for findings in findings_by_device.values()
    for finding in findings
    if finding["severity"] == "Critical"
)

total_high = sum(
    1
    for findings in findings_by_device.values()
    for finding in findings
    if finding["severity"] == "High"
)

total_medium = sum(
    1
    for findings in findings_by_device.values()
    for finding in findings
    if finding["severity"] == "Medium"
)

total_low = sum(
    1
    for findings in findings_by_device.values()
    for finding in findings
    if finding["severity"] == "Low"
)


overall_summary = {
    "devices_assessed": total_devices,
    "devices_with_findings": len(findings_by_device),
    "total_findings": total_findings,
    "severity_counts": {
        "Critical": total_critical,
        "High": total_high,
        "Medium": total_medium,
        "Low": total_low,
    },
}


# ============================================================
# Combine Assessment Data
# ============================================================

report_data = {
    "overall_summary": overall_summary,
    "device_summaries": device_summaries,
    "findings": dict(findings_by_device),
}


# ============================================================
# Save Assessment Data
# ============================================================

with open(
    "assessment.json",
    "w",
    encoding="utf-8"
) as file:

    json.dump(
        report_data,
        file,
        indent=4
    )


# ============================================================
# Gemini Report Generation
# ============================================================

report_prompt = """
You are a cybersecurity risk assessment report writer.

Generate professional report content using ONLY the supplied assessment data.

The Python assessment engine is the authoritative source for:

- device IDs
- findings
- likelihood
- impact
- risk scores
- severity classifications
- NIST CSF mappings
- controls
- recommendations
- assessment totals

Do not modify, recalculate, or reinterpret these values.

Do not invent:

- vulnerabilities
- incidents
- attacks
- compromised devices
- dates
- organizations
- assessors
- asset owners
- compliance certifications
- controls
- findings

You may provide concise explanatory context about why documented findings
matter from an IoT cybersecurity risk perspective.

Keep explanatory statements general and conditional.

Do not assume the presence of sensitive data, critical operations,
known exploitable vulnerabilities, successful attacks, or specific
attack paths unless explicitly provided in the assessment data.

Use phrases such as "may increase risk", "could expose", or
"indicates a security gap" where appropriate.

The report should contain:

1. Executive Summary
2. Overall Assessment
3. Risk Overview
4. Device-Level Findings
5. NIST CSF 2.0 Mapping
6. Recommended Remediation
7. Conclusion

Return the following JSON structure:

{
    "title": "...",
    "executive_summary": "...",
    "overall_assessment": "...",
    "risk_overview": "...",
    "device_findings": [
        {
            "device_id": "...",
            "summary": "..."
        }
    ],
    "nist_mapping": [
        {
            "control_id": "...",
            "explanation": "..."
        }
    ],
    "recommended_remediation": [
        "..."
    ],
    "conclusion": "..."
}

Important:

- Do not calculate totals yourself.
- Do not introduce numerical assessment results
  that are not present in the supplied data.
- Do not change any severity classification.
- Do not claim an incident occurred.
- If information is unavailable, state "Not Provided".
- Return ONLY valid JSON.
"""


client = genai.Client()


request = (
    report_prompt
    + "\n\nAssessment Data:\n"
    + json.dumps(report_data, indent=4)
)


chat = client.chats.create(
    model="gemini-2.5-flash"
)

response = chat.send_message(request)


# ============================================================
# Parse Gemini Response
# ============================================================

response_text = response.text.strip()

if response_text.startswith("```"):
    response_text = response_text.split("\n", 1)[1]
    response_text = response_text.rsplit("```", 1)[0]

try:
    report_content = json.loads(response_text)

except json.JSONDecodeError:
    print("ERROR: Gemini returned invalid JSON.")
    print(response.text)
    raise


# ============================================================
# Save Gemini Report Content
# ============================================================

with open(
    "report_content.json",
    "w",
    encoding="utf-8"
) as file:

    json.dump(
        report_content,
        file,
        indent=4
    )


# ============================================================
# Generate DOCX Report
# ============================================================

document = Document()


# ------------------------------------------------------------
# Title
# ------------------------------------------------------------

document.add_heading(
    report_content["title"],
    level=0
)


# ------------------------------------------------------------
# Executive Summary
# ------------------------------------------------------------

document.add_heading(
    "Executive Summary",
    level=1
)

document.add_paragraph(
    report_content["executive_summary"]
)


# ------------------------------------------------------------
# Overall Assessment
# ------------------------------------------------------------

document.add_heading(
    "Overall Assessment",
    level=1
)

document.add_paragraph(
    report_content["overall_assessment"]
)


# ------------------------------------------------------------
# Assessment Statistics
# ------------------------------------------------------------

document.add_heading(
    "Assessment Statistics",
    level=2
)

summary = overall_summary

document.add_paragraph(
    f"Devices assessed: {summary['devices_assessed']}"
)

document.add_paragraph(
    f"Devices with findings: {summary['devices_with_findings']}"
)

document.add_paragraph(
    f"Total findings: {summary['total_findings']}"
)

document.add_paragraph(
    f"Critical findings: {summary['severity_counts']['Critical']}"
)

document.add_paragraph(
    f"High findings: {summary['severity_counts']['High']}"
)

document.add_paragraph(
    f"Medium findings: {summary['severity_counts']['Medium']}"
)

document.add_paragraph(
    f"Low findings: {summary['severity_counts']['Low']}"
)


# ------------------------------------------------------------
# Risk Overview
# ------------------------------------------------------------

document.add_heading(
    "Risk Overview",
    level=1
)

document.add_paragraph(
    report_content["risk_overview"]
)


# ------------------------------------------------------------
# Device-Level Findings
# ------------------------------------------------------------

document.add_heading(
    "Device-Level Findings",
    level=1
)


for device_id, findings in findings_by_device.items():

    device_summary = device_summaries[device_id]

    document.add_heading(
        device_id,
        level=2
    )

    document.add_paragraph(
        f"Total findings: {device_summary['total_findings']}"
    )

    document.add_paragraph(
        f"Critical: {device_summary['critical']} | "
        f"High: {device_summary['high']} | "
        f"Medium: {device_summary['medium']} | "
        f"Low: {device_summary['low']}"
    )

    highest_risk = device_summary["highest_risk"]

    document.add_paragraph(
        f"Highest-risk finding: "
        f"{highest_risk['issue']} "
        f"(Risk Score: {highest_risk['risk_score']}, "
        f"Severity: {highest_risk['severity']})"
    )

    table = document.add_table(
        rows=1,
        cols=5
    )

    header = table.rows[0].cells

    header[0].text = "Finding"
    header[1].text = "Risk Score"
    header[2].text = "Severity"
    header[3].text = "NIST CSF"
    header[4].text = "Recommendation"

    for finding in findings:

        row = table.add_row().cells

        row[0].text = finding["issue"]
        row[1].text = str(finding["risk_score"])
        row[2].text = finding["severity"]
        row[3].text = finding["csf_control"]
        row[4].text = finding["recommendation"]


# ------------------------------------------------------------
# NIST CSF Mapping
# ------------------------------------------------------------

document.add_heading(
    "NIST CSF 2.0 Mapping",
    level=1
)


for mapping in report_content["nist_mapping"]:

    document.add_heading(
        mapping["control_id"],
        level=2
    )

    document.add_paragraph(
        mapping["explanation"]
    )


# ------------------------------------------------------------
# Recommended Remediation
# ------------------------------------------------------------

document.add_heading(
    "Recommended Remediation",
    level=1
)


for recommendation in report_content["recommended_remediation"]:

    document.add_paragraph(
        recommendation,
        style="List Bullet"
    )


# ------------------------------------------------------------
# Conclusion
# ------------------------------------------------------------

document.add_heading(
    "Conclusion",
    level=1
)

document.add_paragraph(
    report_content["conclusion"]
)


# ============================================================
# Save DOCX
# ============================================================

output_file = "iot_security_assessment_report.docx"

document.save(output_file)

print()
print("=" * 60)
print("IoT SECURITY RISK ASSESSMENT COMPLETE")
print("=" * 60)
print()
print(f"Devices assessed: {total_devices}")
print(f"Devices with findings: {len(findings_by_device)}")
print(f"Total findings: {total_findings}")
print(f"Critical: {total_critical}")
print(f"High: {total_high}")
print(f"Medium: {total_medium}")
print(f"Low: {total_low}")
print()
print(f"Assessment data: assessment.json")
print(f"AI report content: report_content.json")
print(f"Final report: {output_file}")