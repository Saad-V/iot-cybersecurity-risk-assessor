from docx import Document
from main import report_content, overall_summary, device_summaries, findings_by_device, total_devices, total_findings, total_critical, total_high, total_medium, total_low

document = Document()


# ------------------------------------------------------------
# Title
# ------------------------------------------------------------

document.add_heading(
    report_content["title"],
    level=0
)


# ------------------------------------------------------------
# Executive Summary
# ------------------------------------------------------------

document.add_heading(
    "Executive Summary",
    level=1
)

document.add_paragraph(
    report_content["executive_summary"]
)


# ------------------------------------------------------------
# Overall Assessment
# ------------------------------------------------------------

document.add_heading(
    "Overall Assessment",
    level=1
)

document.add_paragraph(
    report_content["overall_assessment"]
)


# ------------------------------------------------------------
# Assessment Statistics
# ------------------------------------------------------------

document.add_heading(
    "Assessment Statistics",
    level=2
)

summary = overall_summary

document.add_paragraph(
    f"Devices assessed: {summary['devices_assessed']}"
)

document.add_paragraph(
    f"Devices with findings: {summary['devices_with_findings']}"
)

document.add_paragraph(
    f"Total findings: {summary['total_findings']}"
)

document.add_paragraph(
    f"Critical findings: {summary['severity_counts']['Critical']}"
)

document.add_paragraph(
    f"High findings: {summary['severity_counts']['High']}"
)

document.add_paragraph(
    f"Medium findings: {summary['severity_counts']['Medium']}"
)

document.add_paragraph(
    f"Low findings: {summary['severity_counts']['Low']}"
)


# ------------------------------------------------------------
# Risk Overview
# ------------------------------------------------------------

document.add_heading(
    "Risk Overview",
    level=1
)

document.add_paragraph(
    report_content["risk_overview"]
)


# ------------------------------------------------------------
# Device-Level Findings
# ------------------------------------------------------------

document.add_heading(
    "Device-Level Findings",
    level=1
)


for device_id, findings in findings_by_device.items():

    device_summary = device_summaries[device_id]

    document.add_heading(
        device_id,
        level=2
    )

    document.add_paragraph(
        f"Total findings: {device_summary['total_findings']}"
    )

    document.add_paragraph(
        f"Critical: {device_summary['critical']} | "
        f"High: {device_summary['high']} | "
        f"Medium: {device_summary['medium']} | "
        f"Low: {device_summary['low']}"
    )

    highest_risk = device_summary["highest_risk"]

    document.add_paragraph(
        f"Highest-risk finding: "
        f"{highest_risk['issue']} "
        f"(Risk Score: {highest_risk['risk_score']}, "
        f"Severity: {highest_risk['severity']})"
    )

    table = document.add_table(
        rows=1,
        cols=5
    )

    header = table.rows[0].cells

    header[0].text = "Finding"
    header[1].text = "Risk Score"
    header[2].text = "Severity"
    header[3].text = "NIST CSF"
    header[4].text = "Recommendation"

    for finding in findings:

        row = table.add_row().cells

        row[0].text = finding["issue"]
        row[1].text = str(finding["risk_score"])
        row[2].text = finding["severity"]
        row[3].text = finding["csf_control"]
        row[4].text = finding["recommendation"]


# ------------------------------------------------------------
# NIST CSF Mapping
# ------------------------------------------------------------

document.add_heading(
    "NIST CSF 2.0 Mapping",
    level=1
)


for mapping in report_content["nist_mapping"]:

    document.add_heading(
        mapping["control_id"],
        level=2
    )

    document.add_paragraph(
        mapping["explanation"]
    )


# ------------------------------------------------------------
# Recommended Remediation
# ------------------------------------------------------------

document.add_heading(
    "Recommended Remediation",
    level=1
)


for recommendation in report_content["recommended_remediation"]:

    document.add_paragraph(
        recommendation,
        style="List Bullet"
    )


# ------------------------------------------------------------
# Conclusion
# ------------------------------------------------------------

document.add_heading(
    "Conclusion",
    level=1
)

document.add_paragraph(
    report_content["conclusion"]
)


# ============================================================
# Save DOCX
# ============================================================

output_file = "iot_security_assessment_report.docx"

document.save(output_file)

print()
print("=" * 60)
print("IoT SECURITY RISK ASSESSMENT COMPLETE")
print("=" * 60)
print()
print(f"Devices assessed: {total_devices}")
print(f"Devices with findings: {len(findings_by_device)}")
print(f"Total findings: {total_findings}")
print(f"Critical: {total_critical}")
print(f"High: {total_high}")
print(f"Medium: {total_medium}")
print(f"Low: {total_low}")
print()
print(f"Assessment data: assessment.json")
print(f"AI report content: report_content.json")
print(f"Final report: {output_file}")